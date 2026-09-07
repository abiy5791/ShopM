"""Product import/export: money units, bad input, the PDF catalogue, and the HTTP surface."""

import io

import pytest
from openpyxl import Workbook, load_workbook

from apps.activity.models import ActivityLog
from apps.catalog import excel, pdf
from apps.catalog.models import Category, Product, Supplier
from apps.inventory.models import InventoryTransaction
from apps.inventory.services import ledger_stock, record_transaction
from apps.purchases.models import Purchase

pytestmark = pytest.mark.django_db

HEADER = excel.COLUMNS


def _sheet(rows, header=HEADER) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    if header is not None:
        ws.append(header)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    buffer.name = "products.xlsx"
    return buffer


def _row(**kw):
    values = {
        "sku": "A-1",
        "name": "Cola",
        "barcode": "",
        "category": "",
        "supplier": "",
        "unit": "pcs",
        "purchase_price": 90,
        "selling_price": 150,
        "opening_stock": 0,
        "min_stock_alert": 5,
        "status": "active",
    }
    values.update(kw)
    return [values[column] for column in HEADER]


# --- money is written in major units, and comes back exact ---
def test_export_writes_major_units(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))
    Product.objects.create(shop=shop, sku="A-1", name="Cola", selling_price=15000)

    ws = load_workbook(
        io.BytesIO(excel.export_products(shop.products.all(), currency="ETB"))
    ).active
    assert [c.value for c in ws[1]] == HEADER
    assert ws.cell(row=2, column=HEADER.index("selling_price") + 1).value == 150


def test_import_reads_major_units(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    result = excel.import_products(
        _sheet([_row(purchase_price="90.55", selling_price=150)]), shop=shop, currency="ETB"
    )

    assert (result["created"], result["skipped"]) == (1, 0)
    product = shop.products.get(sku="A-1")
    assert (product.purchase_price, product.selling_price) == (9055, 15000)


def test_zero_decimal_currency_round_trips(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))
    Product.objects.create(shop=shop, sku="A-1", name="Cola", selling_price=150)

    blob = excel.export_products(shop.products.all(), currency="JPY")
    shop.products.all().delete()
    excel.import_products(io.BytesIO(blob), shop=shop, currency="JPY")

    assert shop.products.get(sku="A-1").selling_price == 150


def test_supplier_round_trips(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))
    supplier = Supplier.objects.create(shop=shop, name="Acme")
    Product.objects.create(shop=shop, sku="A-1", name="Cola", supplier=supplier)
    dst = make_shop(make_user("o2@shopm.local"), name="Dst")

    excel.import_products(
        io.BytesIO(excel.export_products(shop.products.all())), shop=dst, currency="ETB"
    )

    assert dst.products.get(sku="A-1").supplier.name == "Acme"


# --- opening stock: seeded through the ledger, never written directly ---
def test_opening_stock_creates_a_ledger_row(make_user, make_shop):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)

    result = excel.import_products(
        _sheet([_row(opening_stock=12)]), shop=shop, currency="ETB", user=owner
    )

    product = shop.products.get()
    assert result["stock_set"] == 1
    assert product.stock_cached == 12
    assert ledger_stock(product) == 12  # the ledger, not just the cache
    # Booked as a purchase, so the ledger row carries its cost and points back
    # at the purchase the stock came in on.
    txn = InventoryTransaction.objects.get(product=product)
    assert (txn.quantity, txn.type, txn.user) == (12, "purchase", owner)
    assert (txn.unit_cost, txn.reference_type) == (9000, "purchase")


def test_a_quantity_on_an_existing_product_restocks_it(make_user, make_shop):
    """A quantity stocks its row in whether the product is new or not.

    So importing the same file twice does add the stock twice. That is the
    point of the dry run, which reports the quantity every row would add.
    """
    shop = make_shop(make_user("o@shopm.local"))
    sheet = _sheet([_row(opening_stock=12)])
    excel.import_products(sheet, shop=shop, currency="ETB")

    sheet.seek(0)
    result = excel.import_products(sheet, shop=shop, currency="ETB")

    assert (result["updated"], result["stock_set"]) == (1, 1)
    assert shop.products.get().stock_cached == 24
    assert InventoryTransaction.objects.count() == 2


def test_blank_opening_stock_posts_nothing(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    result = excel.import_products(_sheet([_row(opening_stock="")]), shop=shop, currency="ETB")

    assert result["stock_set"] == 0
    assert not InventoryTransaction.objects.exists()


def test_negative_opening_stock_is_a_row_error(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    result = excel.import_products(_sheet([_row(opening_stock=-3)]), shop=shop, currency="ETB")

    assert result["skipped"] == 1
    assert "opening_stock cannot be negative" in result["errors"][0]["error"]
    assert not shop.products.exists()  # the whole row rolled back


def test_export_carries_stock_so_a_catalogue_copy_is_complete(make_user, make_shop):
    src = make_shop(make_user("o@shopm.local"), name="Src")
    product = Product.objects.create(shop=src, sku="A-1", name="Cola", selling_price=15000)
    record_transaction(product=product, quantity=7, type="adjustment")
    dst = make_shop(make_user("o2@shopm.local"), name="Dst")

    excel.import_products(
        io.BytesIO(excel.export_products(src.products.all())), shop=dst, currency="ETB"
    )

    assert dst.products.get(sku="A-1").stock_cached == 7


# --- numeric-looking cells keep their identity ---
def test_numeric_sku_and_barcode_are_not_floats(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    excel.import_products(_sheet([_row(sku=12345, barcode=6001234567890)]), shop=shop)

    product = shop.products.get()
    assert (product.sku, product.barcode) == ("12345", "6001234567890")


def test_prices_tolerate_symbols_and_separators(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    excel.import_products(_sheet([_row(selling_price="Br 1,500.00")]), shop=shop, currency="ETB")

    assert shop.products.get().selling_price == 150000


# --- one bad row is reported and skipped; the rest still land ---
@pytest.mark.parametrize(
    ("bad", "fragment"),
    [
        ({"sku": ""}, "sku is required"),
        ({"selling_price": "abc"}, "not a number"),
        ({"purchase_price": -5}, "cannot be negative"),
        ({"min_stock_alert": -1}, "cannot be negative"),
        ({"status": "archived"}, "status"),
        ({"sku": "X" * 65}, "longer than"),
    ],
)
def test_bad_row_is_skipped_not_fatal(make_user, make_shop, bad, fragment):
    shop = make_shop(make_user("o@shopm.local"))

    result = excel.import_products(
        _sheet([_row(**bad), _row(sku="GOOD-1")]), shop=shop, currency="ETB"
    )

    assert (result["created"], result["skipped"]) == (1, 1)
    assert result["errors"][0]["row"] == 2
    assert fragment in result["errors"][0]["error"]
    assert list(shop.products.values_list("sku", flat=True)) == ["GOOD-1"]


def test_rejected_row_leaves_no_stray_category(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    excel.import_products(
        _sheet([_row(status="archived", category="Ghost")]), shop=shop, currency="ETB"
    )

    assert not Category.objects.filter(shop=shop).exists()


def test_categories_are_reused_case_insensitively(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))
    Category.objects.create(shop=shop, name="Beverages")

    excel.import_products(
        _sheet([_row(sku="A-1", category="beverages"), _row(sku="A-2", category="BEVERAGES")]),
        shop=shop,
    )

    assert Category.objects.filter(shop=shop).count() == 1


def test_blank_rows_are_ignored(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    result = excel.import_products(_sheet([[None] * len(HEADER), _row()]), shop=shop)

    assert (result["created"], result["skipped"], result["errors"]) == (1, 0, [])


# --- unreadable files fail loudly instead of 500-ing ---
def test_row_cap_imports_the_first_rows_and_says_so(make_user, make_shop, monkeypatch):
    shop = make_shop(make_user("o@shopm.local"))
    monkeypatch.setattr(excel, "MAX_ROWS", 2)

    result = excel.import_products(
        _sheet([_row(sku=f"A-{i}") for i in range(4)]), shop=shop, currency="ETB"
    )

    assert (result["created"], result["skipped"]) == (2, 2)
    assert "Import stopped at 2 rows" in result["errors"][-1]["error"]
    assert shop.products.count() == 2


def test_non_xlsx_file_raises_spreadsheet_error(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    with pytest.raises(excel.SpreadsheetError):
        excel.import_products(io.BytesIO(b"not a spreadsheet"), shop=shop)


def test_missing_sku_header_raises_spreadsheet_error(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    with pytest.raises(excel.SpreadsheetError):
        excel.import_products(_sheet([["Cola", 150]], header=["name", "selling_price"]), shop=shop)


def test_template_has_the_expected_header(make_user, make_shop):
    ws = load_workbook(io.BytesIO(excel.import_template(currency="ETB"))).active
    assert [c.value for c in ws[1]] == HEADER


# --- PDF catalogue ---
def test_catalogue_report_totals_the_stock_it_lists(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))
    for sku, cost, stock, alert in [("A-1", 9000, 3, 1), ("A-2", 5000, 2, 5)]:
        Product.objects.create(
            shop=shop,
            sku=sku,
            name=f"Item {sku}",
            purchase_price=cost,
            stock_cached=stock,
            min_stock_alert=alert,
        )

    report = pdf.catalogue_report(shop.products.order_by("sku"), currency="ETB")
    summary = {item["label"]: item["value"] for item in report["summary"]}

    assert summary["Products"] == 2
    assert summary["Stock value"] == 9000 * 3 + 5000 * 2
    assert summary["Low stock"] == 1  # A-2 only: 2 in stock, alert at 5
    assert report["rows"][0][1] == "A-1"


def test_catalogue_pdf_is_a_valid_document(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))
    Product.objects.create(shop=shop, sku="A-1", name="Cola", selling_price=15000)

    content = pdf.export_products_pdf(shop.products.all(), shop_name="Abebe Retail", currency="ETB")

    assert content[:4] == b"%PDF"
    assert content.rstrip().endswith(b"%%EOF")


def test_catalogue_pdf_handles_an_empty_shop(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))
    assert pdf.export_products_pdf(shop.products.all(), currency="ETB")[:4] == b"%PDF"


# --- HTTP surface ---
def test_export_endpoint_returns_xlsx(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)
    Product.objects.create(shop=shop, sku="A-1", name="Cola", selling_price=15000)

    resp = auth(owner, shop).get("/api/v1/products/export")

    assert resp.status_code == 200
    assert resp["Content-Type"].endswith("spreadsheetml.sheet")
    assert "attachment; filename=" in resp["Content-Disposition"]
    assert load_workbook(io.BytesIO(resp.content)).active.max_row == 2


def test_export_endpoint_returns_pdf_when_asked(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)
    Product.objects.create(shop=shop, sku="A-1", name="Cola", selling_price=15000)

    resp = auth(owner, shop).get("/api/v1/products/export?export=pdf")

    assert resp.status_code == 200
    assert resp["Content-Type"] == "application/pdf"
    assert resp["Content-Disposition"].endswith('.pdf"')
    assert resp.content[:4] == b"%PDF"


def test_export_endpoint_defaults_to_xlsx(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)

    resp = auth(owner, shop).get("/api/v1/products/export?export=csv")

    assert resp.status_code == 200
    assert resp["Content-Type"].endswith("spreadsheetml.sheet")


def test_cashier_can_export_both_formats(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    cashier = make_user("c@shopm.local")
    shop = make_shop(owner, cashier=cashier)
    client = auth(cashier, shop)

    assert client.get("/api/v1/products/export").status_code == 200
    assert client.get("/api/v1/products/export?export=pdf").status_code == 200


def test_template_endpoint_returns_xlsx(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)

    resp = auth(owner, shop).get("/api/v1/products/import-template")

    assert resp.status_code == 200
    assert "products-import-template.xlsx" in resp["Content-Disposition"]


def test_import_endpoint_reports_row_errors(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)
    upload = _sheet([_row(sku="A-1"), _row(sku="")])

    resp = auth(owner, shop).post("/api/v1/products/import", {"file": upload}, format="multipart")

    assert resp.status_code == 200
    assert resp.data["created"] == 1
    assert resp.data["skipped"] == 1
    assert resp.data["errors"][0]["row"] == 3


def test_import_endpoint_rejects_a_non_xlsx_upload(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)
    upload = io.BytesIO(b"sku,name\nA-1,Cola\n")
    upload.name = "products.csv"

    resp = auth(owner, shop).post("/api/v1/products/import", {"file": upload}, format="multipart")

    assert resp.status_code == 400
    assert resp.data["code"] == "invalid_file_type"


def test_import_endpoint_rejects_an_unreadable_xlsx(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)
    upload = io.BytesIO(b"not really a workbook")
    upload.name = "products.xlsx"

    resp = auth(owner, shop).post("/api/v1/products/import", {"file": upload}, format="multipart")

    assert resp.status_code == 400
    assert resp.data["code"] == "invalid_spreadsheet"


def test_cashier_cannot_import_or_get_the_template(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    cashier = make_user("c@shopm.local")
    shop = make_shop(owner, cashier=cashier)
    client = auth(cashier, shop)

    assert client.get("/api/v1/products/import-template").status_code == 403
    assert (
        client.post(
            "/api/v1/products/import", {"file": _sheet([_row()])}, format="multipart"
        ).status_code
        == 403
    )


# --- dry run: plan the import, write nothing, name every collision ---
def test_dry_run_writes_nothing_and_reports_the_overwrite(make_user, make_shop):
    """The production bug: a new sheet reusing a live SKU silently replaced it."""
    shop = make_shop(make_user("o@shopm.local"))
    Product.objects.create(
        shop=shop,
        sku="A-1",
        name="Cola 500ml",
        purchase_price=9000,
        selling_price=15000,
        category=Category.objects.create(shop=shop, name="Beverages"),
    )

    result = excel.import_products(
        _sheet([_row(sku="A-1", name="Notebook A4", selling_price=60, category="Stationery")]),
        shop=shop,
        currency="ETB",
        dry_run=True,
    )

    assert (result["dry_run"], result["created"], result["updated"]) == (True, 0, 1)
    assert result["conflict_count"] == 1
    [conflict] = result["conflicts"]
    assert (conflict["row"], conflict["sku"]) == (2, "A-1")
    assert conflict["existing_name"] == "Cola 500ml"
    assert conflict["duplicate_of_row"] is None
    changes = {c["field"]: (c["from"], c["to"]) for c in conflict["changes"]}
    assert changes["name"] == ("Cola 500ml", "Notebook A4")
    # Money is diffed in the major units the sheet shows, not stored minor units.
    assert changes["selling_price"] == ("150.00", "60.00")
    assert changes["category"] == ("Beverages", "Stationery")

    shop.products.get(sku="A-1").refresh_from_db()
    assert shop.products.get(sku="A-1").name == "Cola 500ml"
    # A planned category is not created either — a preview leaves no trace.
    assert not Category.objects.filter(shop=shop, name="Stationery").exists()


def test_dry_run_reports_no_conflict_for_new_skus(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    result = excel.import_products(
        _sheet([_row(sku="A-1", opening_stock=12), _row(sku="A-2")]),
        shop=shop,
        currency="ETB",
        dry_run=True,
    )

    assert (result["created"], result["updated"], result["stock_set"]) == (2, 0, 1)
    assert result["conflicts"] == []
    assert not shop.products.exists()


def test_dry_run_flags_a_sku_repeated_inside_the_sheet(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    result = excel.import_products(
        _sheet([_row(sku="A-1", name="Cola"), _row(sku="A-1", name="Pen")]),
        shop=shop,
        currency="ETB",
        dry_run=True,
    )

    assert (result["created"], result["updated"]) == (1, 1)
    [conflict] = result["conflicts"]
    # Points back at the earlier row, so the user can fix the file itself.
    assert (conflict["row"], conflict["duplicate_of_row"]) == (3, 2)
    assert conflict["existing_name"] == "Cola"


def test_dry_run_still_reports_row_errors(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    result = excel.import_products(
        _sheet([_row(sku=""), _row(sku="A-1", selling_price="abc")]),
        shop=shop,
        currency="ETB",
        dry_run=True,
    )

    assert result["skipped"] == 2
    assert [e["row"] for e in result["errors"]] == [2, 3]


def test_dry_run_endpoint_leaves_the_catalogue_alone(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)
    Product.objects.create(shop=shop, sku="A-1", name="Cola", selling_price=15000)

    resp = auth(owner, shop).post(
        "/api/v1/products/import?dry_run=1",
        {"file": _sheet([_row(sku="A-1", name="Pen")])},
        format="multipart",
    )

    assert resp.status_code == 200
    assert resp.data["dry_run"] is True
    assert resp.data["conflicts"][0]["existing_name"] == "Cola"
    assert shop.products.get(sku="A-1").name == "Cola"
    # A preview is not an event worth auditing.
    assert not ActivityLog.objects.filter(action="product.import").exists()


def test_import_after_preview_commits(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)
    Product.objects.create(shop=shop, sku="A-1", name="Cola", selling_price=15000)
    client = auth(owner, shop)

    client.post(
        "/api/v1/products/import?dry_run=1",
        {"file": _sheet([_row(sku="A-1", name="Pen")])},
        format="multipart",
    )
    resp = client.post(
        "/api/v1/products/import",
        {"file": _sheet([_row(sku="A-1", name="Pen")])},
        format="multipart",
    )

    assert (resp.status_code, resp.data["updated"]) == (200, 1)
    assert shop.products.get(sku="A-1").name == "Pen"
    assert ActivityLog.objects.filter(action="product.import").count() == 1


def test_dry_run_ignores_a_reimport_that_changes_nothing(make_user, make_shop):
    """Re-importing an untouched export matches every SKU but overwrites nothing."""
    shop = make_shop(make_user("o@shopm.local"))
    excel.import_products(_sheet([_row(sku="A-1")]), shop=shop, currency="ETB")

    result = excel.import_products(
        _sheet([_row(sku="A-1")]), shop=shop, currency="ETB", dry_run=True
    )

    assert (result["updated"], result["conflict_count"]) == (1, 0)
    assert result["conflicts"] == []


# --- imported stock is booked as one purchase, so it has provenance ---
def test_import_books_one_paid_purchase_for_the_whole_sheet(make_user, make_shop):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)
    # Already in the catalogue, so this row is a restock rather than a creation —
    # it still belongs on the purchase.
    Product.objects.create(shop=shop, sku="A-1", name="Cola")

    result = excel.import_products(
        _sheet(
            [
                _row(sku="A-1", supplier="Acme", purchase_price=90, opening_stock=10),
                _row(sku="A-2", supplier="Acme", purchase_price=50, opening_stock=4),
                _row(sku="A-3", supplier="Globex", purchase_price=20, opening_stock=3),
                _row(sku="A-4", supplier="Globex", opening_stock=0),  # nothing to stock in
            ]
        ),
        shop=shop,
        currency="ETB",
        user=owner,
    )

    # One purchase for the run, not one per supplier and not one per row.
    assert (result["stock_set"], result["purchases"]) == (3, 1)
    purchase = Purchase.objects.get()
    # Total is the sum of unit cost x quantity over every line.
    assert purchase.total == 90_00 * 10 + 50_00 * 4 + 20_00 * 3
    # Paid in full, so the import invents no debt.
    assert (purchase.amount_paid, purchase.payment_status) == (purchase.total, "paid")
    assert purchase.notes == excel.IMPORT_PURCHASE_NOTE
    assert {(i.product.sku, i.quantity, i.unit_cost) for i in purchase.items.all()} == {
        ("A-1", 10, 9000),
        ("A-2", 4, 5000),
        ("A-3", 3, 2000),
    }
    # The sheet names two suppliers, so crediting either would be wrong.
    assert purchase.supplier is None
    assert shop.products.get(sku="A-4").stock_cached == 0


def test_a_single_supplier_sheet_credits_that_supplier(make_user, make_shop):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)

    excel.import_products(
        _sheet(
            [
                _row(sku="A-1", supplier="Acme", purchase_price=90, opening_stock=10),
                _row(sku="A-2", supplier="Acme", purchase_price=50, opening_stock=4),
            ]
        ),
        shop=shop,
        currency="ETB",
        user=owner,
    )

    purchase = Purchase.objects.get()
    assert purchase.supplier == Supplier.objects.get(name="Acme")
    # Paid in full, so the supplier is owed nothing for it.
    assert Supplier.objects.get(name="Acme").payable_cached == 0


def test_import_without_a_supplier_still_books_a_purchase(make_user, make_shop):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)

    excel.import_products(
        _sheet([_row(supplier="", opening_stock=5)]), shop=shop, currency="ETB", user=owner
    )

    purchase = Purchase.objects.get()
    assert purchase.supplier is None
    assert ledger_stock(shop.products.get()) == 5


def test_reimport_books_a_second_purchase(make_user, make_shop):
    """Purchase history follows the stock: a second restock is a second purchase."""
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)
    excel.import_products(_sheet([_row(opening_stock=12)]), shop=shop, currency="ETB", user=owner)

    result = excel.import_products(
        _sheet([_row(opening_stock=12)]), shop=shop, currency="ETB", user=owner
    )

    assert (result["updated"], result["stock_set"], result["purchases"]) == (1, 1, 1)
    assert Purchase.objects.count() == 2
    assert shop.products.get().stock_cached == 24


def test_a_sheet_with_no_quantities_books_no_purchase(make_user, make_shop):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)

    result = excel.import_products(
        _sheet([_row(sku="A-1"), _row(sku="A-2")]), shop=shop, currency="ETB", user=owner
    )

    assert (result["created"], result["stock_set"], result["purchases"]) == (2, 0, 0)
    assert not Purchase.objects.exists()


def test_dry_run_books_no_purchase(make_user, make_shop):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)

    excel.import_products(
        _sheet([_row(opening_stock=12)]), shop=shop, currency="ETB", user=owner, dry_run=True
    )

    assert not Purchase.objects.exists()
    assert not InventoryTransaction.objects.exists()


def test_opening_stock_falls_back_to_an_adjustment_without_a_user(make_user, make_shop):
    """Purchase.user is non-null, so a caller with no user still gets the stock."""
    shop = make_shop(make_user("o@shopm.local"))

    result = excel.import_products(_sheet([_row(opening_stock=7)]), shop=shop, currency="ETB")

    assert (result["stock_set"], result["purchases"]) == (1, 0)
    assert not Purchase.objects.exists()
    txn = InventoryTransaction.objects.get()
    assert (txn.type, txn.quantity, txn.notes) == ("adjustment", 7, "Opening stock")


def test_a_rejected_row_books_no_opening_stock(make_user, make_shop):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner)

    result = excel.import_products(
        _sheet([_row(sku="A-1", opening_stock=5), _row(sku="A-2", selling_price="abc")]),
        shop=shop,
        currency="ETB",
        user=owner,
    )

    assert (result["created"], result["skipped"], result["stock_set"]) == (1, 1, 1)
    [item] = Purchase.objects.get().items.all()
    assert item.product.sku == "A-1"


def test_dry_run_lists_the_products_it_would_add(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))

    result = excel.import_products(
        _sheet(
            [
                _row(sku="A-1", name="Cola", supplier="Acme", purchase_price=90, opening_stock=12),
                _row(sku=""),  # rejected, so it is not something we would add
            ]
        ),
        shop=shop,
        currency="ETB",
        dry_run=True,
    )

    assert (result["created"], result["skipped"]) == (1, 1)
    [creation] = result["creations"]
    assert (creation["row"], creation["sku"], creation["name"]) == (2, "A-1", "Cola")
    assert (creation["supplier"], creation["opening_stock"]) == ("Acme", 12)
    # Money in the units the sheet shows, matching the conflict diff.
    assert (creation["purchase_price"], creation["selling_price"]) == ("90.00", "150.00")


def test_dry_run_does_not_list_an_existing_sku_as_a_creation(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))
    Product.objects.create(shop=shop, sku="A-1", name="Cola")

    result = excel.import_products(
        _sheet([_row(sku="A-1", name="Pen"), _row(sku="A-2")]), shop=shop, dry_run=True
    )

    assert [c["sku"] for c in result["creations"]] == ["A-2"]
    assert [c["sku"] for c in result["conflicts"]] == ["A-1"]


def test_dry_run_reports_the_stock_a_row_would_add_to_an_existing_product(make_user, make_shop):
    """The re-import guard is now the preview, so it has to show the quantity."""
    shop = make_shop(make_user("o@shopm.local"))
    excel.import_products(_sheet([_row(sku="A-1", opening_stock=12)]), shop=shop, currency="ETB")

    result = excel.import_products(
        _sheet([_row(sku="A-1", opening_stock=12)]), shop=shop, currency="ETB", dry_run=True
    )

    # Nothing about the product changes, but 12 more units would land.
    assert (result["conflict_count"], result["overwrite_count"]) == (1, 0)
    [conflict] = result["conflicts"]
    assert (conflict["opening_stock"], conflict["changes"]) == (12, [])
    assert result["stock_set"] == 1


def test_dry_run_separates_overwrites_from_plain_restocks(make_user, make_shop):
    shop = make_shop(make_user("o@shopm.local"))
    # Matching _row()'s other columns, so the only differences are the ones under test.
    existing = {"purchase_price": 9000, "selling_price": 15000, "min_stock_alert": 5}
    Product.objects.create(shop=shop, sku="A-1", name="Cola", **existing)
    Product.objects.create(shop=shop, sku="A-2", name="Pen", **existing)

    result = excel.import_products(
        _sheet(
            [
                _row(sku="A-1", name="Notebook", opening_stock=0),  # renames, adds nothing
                _row(sku="A-2", name="Pen", opening_stock=6),  # adds stock, renames nothing
            ]
        ),
        shop=shop,
        currency="ETB",
        dry_run=True,
    )

    assert (result["conflict_count"], result["overwrite_count"]) == (2, 1)
    by_sku = {c["sku"]: c for c in result["conflicts"]}
    assert [c["field"] for c in by_sku["A-1"]["changes"]] == ["name"]
    assert by_sku["A-1"]["opening_stock"] == 0
    assert by_sku["A-2"]["changes"] == []
    assert by_sku["A-2"]["opening_stock"] == 6
