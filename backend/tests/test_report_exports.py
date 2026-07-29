"""Report exports: numbers stay numbers in Excel, and files identify themselves."""

import io

import pytest
from openpyxl import load_workbook

from apps.catalog.models import Product
from apps.reports import exporters

pytestmark = pytest.mark.django_db

REPORT = {
    "key": "sales",
    "title": "Sales report (daily)",
    "currency": "ETB",
    "period": {"start": "2026-07-01", "end": "2026-07-28", "granularity": "daily"},
    "summary": [
        {"label": "Total sales", "value": 125000, "money": True},
        {"label": "Transactions", "value": 7, "money": False},
    ],
    "columns": ["Period", "Transactions", "Total"],
    "rows": [["Hamle 20, 2018", 3, 45000], ["Hamle 21, 2018", 4, 80000]],
    "money_columns": [2],
}


def _sheet(report=REPORT, **kw):
    return load_workbook(io.BytesIO(exporters.to_xlsx(report, **kw))).active


def _column(ws, header: str, report=REPORT):
    """The cells under a table header, wherever the title block pushed the table to.

    Matches the whole header row, not a single cell — summary labels reuse the
    same words as the column names.
    """
    width = len(report["columns"])
    for row in ws.iter_rows():
        if [c.value for c in row[:width]] == report["columns"]:
            column = report["columns"].index(header) + 1
            return [ws.cell(row=r, column=column) for r in range(row[0].row + 1, ws.max_row + 1)]
    raise AssertionError("no table header row")


# --- Excel: money is a number, so an owner can sum the column ---
def test_money_cells_are_numeric_with_a_currency_format():
    total = _column(_sheet(), "Total")[0]

    assert total.value == 450  # 45000 minor units -> 450.00 ETB
    assert not isinstance(total.value, str)
    assert "Br" in total.number_format


def test_counts_stay_numeric_too():
    assert _column(_sheet(), "Transactions")[0].value == 3


def test_zero_decimal_currency_keeps_whole_units():
    report = {**REPORT, "currency": "JPY"}
    assert _column(_sheet(report), "Total")[0].value == 45000


def test_header_block_names_the_shop_and_period():
    ws = _sheet(shop_name="Abebe Retail")
    text = "\n".join(str(c.value) for row in ws.iter_rows(max_row=6) for c in row if c.value)

    assert "Abebe Retail" in text
    assert "Sales report (daily)" in text
    assert "2026-07-01" in text and "Hamle" in text  # both calendars
    assert "ETB" in text


def test_table_is_frozen_and_filterable():
    ws = _sheet()
    assert ws.freeze_panes is not None
    assert ws.auto_filter.ref is not None


def test_empty_report_still_exports():
    ws = _sheet({**REPORT, "rows": []})
    assert ws.auto_filter.ref is None  # nothing to filter, but the file is valid


# --- PDF ---
def test_pdf_is_valid_and_paginated():
    content = exporters.to_pdf(REPORT, shop_name="Abebe Retail")

    assert content[:4] == b"%PDF"
    assert content.rstrip().endswith(b"%%EOF")
    assert b"/Page" in content


def test_pdf_handles_an_empty_report():
    assert exporters.to_pdf({**REPORT, "rows": []}, shop_name="S")[:4] == b"%PDF"


# --- filenames ---
def test_filename_carries_the_date_range():
    assert exporters.filename_for(REPORT, "xlsx") == "sales-2026-07-01_2026-07-28.xlsx"


def test_filename_falls_back_to_today_without_a_period():
    name = exporters.filename_for({"key": "inventory", "period": None}, "pdf")
    assert name.startswith("inventory-") and name.endswith(".pdf")


# --- HTTP surface ---
@pytest.mark.parametrize("fmt", ["pdf", "xlsx"])
def test_export_response_is_named_after_the_report(make_user, make_shop, auth, fmt):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner, name="S")
    Product.objects.create(shop=shop, sku="A-1", name="Cola", selling_price=15000)

    resp = auth(owner, shop).get(f"/api/v1/reports/inventory?export={fmt}")

    assert resp.status_code == 200
    assert 'filename="inventory-' in resp["Content-Disposition"]
    assert resp["Content-Disposition"].endswith(f'.{fmt}"')


def test_unknown_export_format_falls_back_to_json(make_user, make_shop, auth):
    owner = make_user("o@shopm.local")
    shop = make_shop(owner, name="S")

    resp = auth(owner, shop).get("/api/v1/reports/inventory?export=csv")

    assert resp.status_code == 200
    assert resp["Content-Type"].startswith("application/json")
