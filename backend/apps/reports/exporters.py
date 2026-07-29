"""Render a report dict (from services) to XLSX or PDF (plan §11 Phase 5).

Both formats open with the same header block — shop, report, period, generated-at
— so a file still explains itself once it has left the app.

Money in the .xlsx is written as a **number** with a currency cell format, never
as pre-formatted text: an owner who exports to Excel expects to sum, sort and
pivot the column. The PDF, being a finished document, prints formatted strings.
"""

import io
import re
from datetime import date

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.common.ethiopian import format_ethiopian
from apps.common.money import CURRENCY_SYMBOLS, exponent, format_money, to_major

HEADER_FILL = "334155"  # slate-700, the app's table header
BORDER = colors.HexColor("#E6E8EA")
ZEBRA = colors.HexColor("#F8FAFC")
MUTED = colors.HexColor("#64748B")

PAGE_MARGIN = 15 * mm
#: Sheet names can't contain these, and Excel silently corrupts files that try.
_SHEET_TITLE_JUNK = re.compile(r"[\[\]:*?/\\]")


# --------------------------------------------------------------------------
# Shared header text
# --------------------------------------------------------------------------
def _period_line(report: dict) -> str:
    """ "Hamle 20, 2018 – Hamle 21, 2018 (2026-07-27 – 2026-07-28)", or "" if the
    report has no date range (inventory is a point-in-time snapshot)."""
    period = report.get("period") or {}
    start, end = period.get("start"), period.get("end")
    if not (start and end):
        return ""
    try:
        first, last = date.fromisoformat(start), date.fromisoformat(end)
    except (TypeError, ValueError):
        return f"{start} – {end}"
    return f"{format_ethiopian(first)} – {format_ethiopian(last)}  ({start} – {end})"


def _generated_line() -> str:
    return f"Generated {timezone.localtime():%Y-%m-%d %H:%M}"


def _header_lines(report: dict) -> list[str]:
    """Sub-title lines: period, currency, timestamp — skipping what doesn't apply."""
    lines = []
    period = _period_line(report)
    if period:
        lines.append(f"Period: {period}")
    lines.append(f"Amounts in {report['currency'].upper()} · {_generated_line()}")
    return lines


def filename_for(report: dict, fmt: str) -> str:
    """e.g. "sales-2026-06-28_2026-07-28.xlsx" — a date-stamped name so a folder
    of exports stays sortable and nothing silently overwrites."""
    period = report.get("period") or {}
    start, end = period.get("start"), period.get("end")
    stamp = f"{start}_{end}" if start and end else f"{timezone.localdate():%Y-%m-%d}"
    return f"{report['key']}-{stamp}.{fmt}"


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------
def _excel_money_format(currency: str) -> str:
    """A number format that shows "Br1,250.00" while the cell stays numeric."""
    decimals = exponent(currency)
    number = "#,##0" + ("." + "0" * decimals if decimals else "")
    symbol = CURRENCY_SYMBOLS.get(currency.upper())
    return f'"{symbol}"{number}' if symbol else f'{number} "{currency.upper()}"'


def _is_number(value) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool)


def _write(ws, row: int, column: int, value, *, money: bool, money_format: str, currency: str):
    cell = ws.cell(row=row, column=column)
    if money and _is_number(value):
        cell.value = to_major(int(value), currency)
        cell.number_format = money_format
    elif _is_number(value):
        cell.value = value
        cell.number_format = "#,##0"
    else:
        cell.value = "" if value is None else str(value)
    return cell


def to_xlsx(report: dict, *, shop_name: str = "") -> bytes:
    currency = report["currency"]
    money_format = _excel_money_format(currency)
    money_cols = set(report.get("money_columns", []))
    columns = report["columns"]

    wb = Workbook()
    ws = wb.active
    ws.title = _SHEET_TITLE_JUNK.sub("", report["title"])[:31] or "Report"

    if shop_name:
        ws.append([shop_name])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=14)
    ws.append([report["title"]])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    for line in _header_lines(report):
        ws.append([line])
        ws.cell(row=ws.max_row, column=1).font = Font(size=9, color="FF64748B")

    ws.append([])
    for item in report["summary"]:
        row = ws.max_row + 1
        ws.cell(row=row, column=1, value=item["label"]).font = Font(bold=True)
        _write(
            ws,
            row,
            2,
            item["value"],
            money=bool(item.get("money")),
            money_format=money_format,
            currency=currency,
        )

    ws.append([])
    header_row = ws.max_row + 1
    for index, name in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=index, value=name)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center")

    for offset, values in enumerate(report["rows"], start=1):
        for index, value in enumerate(values):
            _write(
                ws,
                header_row + offset,
                index + 1,
                value,
                money=index in money_cols,
                money_format=money_format,
                currency=currency,
            )

    last_row = header_row + len(report["rows"])
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if report["rows"]:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{last_row}"

    # Width from the widest cell in each column, so nothing opens as "#####".
    for index, name in enumerate(columns, start=1):
        longest = max(
            [len(str(name))]
            + [len(str(row[index - 1])) for row in report["rows"] if index <= len(row)]
        )
        ws.column_dimensions[get_column_letter(index)].width = min(max(longest + 4, 12), 48)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
class _NumberedCanvas(pdf_canvas.Canvas):
    """Buffers pages so the footer can say "Page 1 of 4" — the total isn't known
    until the document is fully laid out."""

    def __init__(self, *args, footer_text="", **kwargs):
        super().__init__(*args, **kwargs)
        self._footer_text = footer_text
        self._pages: list[dict] = []

    def showPage(self):  # noqa: N802 (reportlab API)
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._pages)
        for state in self._pages:
            self.__dict__.update(state)
            self._draw_footer(total)
            super().showPage()
        super().save()

    def _draw_footer(self, total: int):
        self.setFont("Helvetica", 7)
        self.setFillColor(MUTED)
        self.drawString(PAGE_MARGIN, 10 * mm, self._footer_text)
        self.drawRightString(A4[0] - PAGE_MARGIN, 10 * mm, f"Page {self._pageNumber} of {total}")


def _pdf_cell(value, index: int, money_cols: set, currency: str) -> str:
    if index in money_cols and _is_number(value):
        return format_money(int(value), currency)
    if _is_number(value):
        return f"{value:,}"
    return "" if value is None else str(value)


def _column_widths(table_data: list[list[str]], available: float) -> list[float]:
    """Share the frame width out by how much text each column actually holds."""
    count = len(table_data[0])
    weights = [max(len(row[i]) for row in table_data) or 1 for i in range(count)]
    total = sum(weights)
    widths = [max(available * w / total, 18 * mm) for w in weights]
    overflow = sum(widths) / available
    return [w / overflow for w in widths] if overflow > 1 else widths


def _wrap(text: str, style: ParagraphStyle, threshold: int = 28):
    """Long labels (product names) wrap; short cells stay plain for speed."""
    return Paragraph(text, style) if len(text) > threshold else text


def to_pdf(report: dict, *, shop_name: str = "") -> bytes:
    buffer = io.BytesIO()
    currency = report["currency"]
    money_cols = set(report.get("money_columns", []))
    styles = getSampleStyleSheet()
    body = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=8, leading=10, spaceAfter=0)
    sub = ParagraphStyle(
        "sub", parent=styles["BodyText"], fontSize=8.5, textColor=MUTED, alignment=TA_LEFT
    )

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        title=f"{report['title']} — {shop_name}" if shop_name else report["title"],
        author=shop_name or "ShopM",
        leftMargin=PAGE_MARGIN,
        rightMargin=PAGE_MARGIN,
        topMargin=PAGE_MARGIN,
        bottomMargin=18 * mm,
    )
    available = doc.width

    elements = []
    if shop_name:
        elements.append(Paragraph(shop_name, styles["Title"]))
    elements.append(Paragraph(report["title"], styles["Heading2" if shop_name else "Title"]))
    for line in _header_lines(report):
        elements.append(Paragraph(line, sub))
    elements.append(Spacer(1, 6 * mm))

    summary_rows = [
        [item["label"], _pdf_cell(item["value"], 1, {1} if item.get("money") else set(), currency)]
        for item in report["summary"]
    ]
    if summary_rows:
        summary = Table(summary_rows, hAlign="LEFT", colWidths=[available * 0.35, available * 0.2])
        summary.setStyle(
            TableStyle(
                [
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("TEXTCOLOR", (0, 0), (0, -1), MUTED),
                    ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
                    ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        elements += [summary, Spacer(1, 6 * mm)]

    if report["rows"]:
        text_rows = [list(report["columns"])] + [
            [_pdf_cell(v, i, money_cols, currency) for i, v in enumerate(row)]
            for row in report["rows"]
        ]
        widths = _column_widths(text_rows, available)
        table_data = [text_rows[0]] + [[_wrap(cell, body) for cell in row] for row in text_rows[1:]]
        table = Table(table_data, repeatRows=1, colWidths=widths)
        numeric = sorted(money_cols | {i for i, v in enumerate(report["rows"][0]) if _is_number(v)})
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{HEADER_FILL}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.25, BORDER),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ZEBRA]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        # Figures right-align; only text reads well flush left.
        style += [("ALIGN", (i, 0), (i, -1), "RIGHT") for i in numeric]
        table.setStyle(TableStyle(style))
        elements.append(table)
    else:
        elements.append(Paragraph("No data for this period.", sub))

    footer = " · ".join(filter(None, [shop_name, report["title"], _generated_line()]))
    doc.build(
        elements,
        canvasmaker=lambda *args, **kwargs: _NumberedCanvas(*args, footer_text=footer, **kwargs),
    )
    return buffer.getvalue()
